#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Add a formatted per-application Summary sheet to an exported Azure estimate.

    python build_summary.py "ExportedEstimate.xlsx" --out "Azure-BoQ-v1.0.xlsx"

Reads the exported sheet (header on row 3, data from row 4 to the Total row) and
writes a Summary sheet with a KPI block plus cost by application, region and
pricing model.

Application = the first '|'-delimited segment of the Custom name, which is why
line items should be named "<App> | <Server> | <SKU> | <Disk>".

Values are written as STATIC numbers, not cross-sheet formulas: the app's Excel
preview renders cross-sheet SUMIF as #ERROR!, and a BoQ is a point-in-time quote.
"""
import argparse, re, shutil, sys
from collections import OrderedDict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl required:  pip install openpyxl")

NAVY, BLUE, LIGHT = "1F3864", "2E5B9A", "DCE6F1"
MONEY, PCT = '#,##0.00', '0.0%'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--out")
    ap.add_argument("--sheet", help="BoQ sheet name (default: the non-Summary sheet)")
    ap.add_argument("--title", default="Azure Bill of Quantities")
    a = ap.parse_args()

    out = a.out or a.workbook
    if out != a.workbook:
        shutil.copyfile(a.workbook, out)
    wb = openpyxl.load_workbook(out)

    name = a.sheet or next((s for s in wb.sheetnames if s.lower() != "summary"), None)
    boq = wb[name]

    # locate header row, then read until the Total row
    hdr = next((i for i in range(1, 8)
                if any("Custom name" == str(c or "") for c in
                       next(boq.iter_rows(min_row=i, max_row=i, values_only=True)))), 3)
    cols = {str(c or ""): j for j, c in
            enumerate(next(boq.iter_rows(min_row=hdr, max_row=hdr, values_only=True)))}
    cai = cols.get("Service category", 0)
    si = cols.get("Service type", 1)
    ci = cols.get("Custom name", 2)
    ri = cols.get("Region", 3)
    di = cols.get("Description", 4)
    mi = cols.get("Estimated monthly cost", 5)

    apps, regions, models, cats = OrderedDict(), OrderedDict(), OrderedDict(), OrderedDict()
    total, nrows = 0.0, 0
    for r in boq.iter_rows(min_row=hdr + 1, values_only=True):
        if not r:
            continue
        cell = lambda j: str(r[j] or "").strip() if j is not None and j < len(r) else ""
        # the Total row puts its label in the Region column, not Custom name
        if any(cell(j).lower() == "total" for j in (ci, ri, di) if j is not None):
            break
        svc = cell(si)
        if not svc:                       # Support / Licensing / Billing metadata rows
            continue
        try:
            cost = float(r[mi] or 0)
        except (TypeError, ValueError):
            continue

        nm, desc, reg, cat = cell(ci), cell(di), cell(ri), cell(cai)
        # shared services often have no custom name — group them by service type
        app = nm.split("|")[0].strip() if nm else svc
        model = ("3-Year Reserved Instance" if "3 year reserved" in desc
                 else "1-Year Reserved Instance" if "1 year reserved" in desc
                 else "3-Year Savings Plan" if "3 year savings" in desc
                 else "1-Year Savings Plan" if "1 year savings" in desc
                 else "Pay as you go")
        d = apps.setdefault(app, {"n": 0, "cost": 0.0, "regions": set(), "vcpu": 0})
        d["n"] += 1; d["cost"] += cost; d["regions"].add(reg)
        m = re.search(r"\((\d+) vCPUs", desc)
        if m:
            d["vcpu"] += int(m.group(1))
        for bucket, key in ((regions, reg), (models, model), (cats, cat or svc)):
            b = bucket.setdefault(key, {"n": 0, "cost": 0.0})
            b["n"] += 1; b["cost"] += cost
        total += cost; nrows += 1

    if not nrows:
        sys.exit("No priced line items found — check --sheet.")

    order = sorted(apps.items(), key=lambda kv: -kv[1]["cost"])
    # sum the rounded components so every displayed figure ties exactly
    disp_month = round(sum(round(d["cost"], 2) for _, d in order), 2)
    disp_year = round(sum(round(d["cost"] * 12, 2) for _, d in order), 2)

    if "Summary" in wb.sheetnames:
        wb.remove(wb["Summary"])
    ws = wb.create_sheet("Summary", 0)

    thin = Side(style="thin", color="B4C6E7")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def put(cell, val, *, bold=False, size=11, color=None, fill=None,
            fmt=None, align=None, border=True, wrap=False):
        c = ws[cell]; c.value = val
        c.font = Font(bold=bold, size=size, color=color or "000000")
        if fill:  c.fill = PatternFill("solid", fgColor=fill)
        if fmt:   c.number_format = fmt
        if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if border: c.border = box

    def band(row, text, span):
        ws.merge_cells(f"A{row}:{span}{row}")
        put(f"A{row}", text, bold=True, size=12, color="FFFFFF", fill=NAVY,
            align="left", border=False)

    def head(row, labels):
        for i, lab in enumerate(labels):
            put(f"{get_column_letter(i+1)}{row}", lab, bold=True, color="FFFFFF",
                fill=BLUE, align="center" if i else "left", wrap=True)
        ws.row_dimensions[row].height = 30

    ws.merge_cells("A1:G1")
    put("A1", a.title, bold=True, size=16, color="FFFFFF", fill=NAVY,
        align="left", border=False)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:G2")
    put("A2", f"Cost summary by application  ·  {nrows} line items  ·  "
              f"{', '.join(sorted(regions))}  ·  USD",
        size=10, color="444444", align="left", border=False)

    for i, (lab, val, fmt) in enumerate([
            ("Total line items", nrows, '#,##0'),
            ("Monthly Cost", disp_month, MONEY),
            ("Annual Cost", disp_year, MONEY),
            ("3-Year Cost", round(disp_year * 3, 2), MONEY)]):
        put(f"A{4+i}", lab, bold=True, fill=LIGHT)
        put(f"B{4+i}", val, bold=True, fmt=fmt, align="right")

    T1 = 10
    band(T1 - 1, "COST BY APPLICATION", "G")
    head(T1, ["Application", "Items", "vCPUs", "Region", "Monthly (USD)",
              "Annual (USD)", "% of Total"])
    r = T1 + 1
    for app, d in order:
        put(f"A{r}", app, bold=True)
        put(f"B{r}", d["n"], fmt='#,##0', align="center")
        put(f"C{r}", d["vcpu"], fmt='#,##0', align="center")
        put(f"D{r}", ", ".join(sorted(x for x in d["regions"] if x)), align="center")
        put(f"E{r}", round(d["cost"], 2), fmt=MONEY, align="right")
        put(f"F{r}", round(d["cost"] * 12, 2), fmt=MONEY, align="right")
        put(f"G{r}", d["cost"] / total, fmt=PCT, align="center")
        r += 1
    for col, f in (("A", "TOTAL"), ("B", f"=SUM(B{T1+1}:B{r-1})"),
                   ("C", f"=SUM(C{T1+1}:C{r-1})"), ("D", ""),
                   ("E", f"=SUM(E{T1+1}:E{r-1})"), ("F", f"=SUM(F{T1+1}:F{r-1})"),
                   ("G", f"=SUM(G{T1+1}:G{r-1})")):
        fmt = MONEY if col in "EF" else (PCT if col == "G" else '#,##0')
        put(f"{col}{r}", f, bold=True, color="FFFFFF", fill=NAVY,
            fmt=None if col in "AD" else fmt,
            align="right" if col in "EF" else "center" if col != "A" else "left")

    def small_table(start, title, data):
        band(start - 1, title, "C")
        head(start, [title.replace("COST BY ", "").title(), "Items", "Monthly (USD)"])
        rr = start + 1
        for k, d in sorted(data.items(), key=lambda kv: -kv[1]["cost"]):
            put(f"A{rr}", k or "(unspecified)")
            put(f"B{rr}", d["n"], fmt='#,##0', align="center")
            put(f"C{rr}", round(d["cost"], 2), fmt=MONEY, align="right")
            rr += 1
        put(f"A{rr}", "TOTAL", bold=True, fill=LIGHT)
        put(f"B{rr}", f"=SUM(B{start+1}:B{rr-1})", bold=True, fill=LIGHT, fmt='#,##0', align="center")
        put(f"C{rr}", f"=SUM(C{start+1}:C{rr-1})", bold=True, fill=LIGHT, fmt=MONEY, align="right")
        return rr

    r2 = small_table(r + 3, "COST BY SERVICE CATEGORY", cats)
    r3 = small_table(r2 + 3, "COST BY REGION", regions)
    r4 = small_table(r3 + 3, "COST BY PRICING MODEL", models)

    n = r4 + 2
    put(f"A{n}", "Notes & assumptions", bold=True, border=False)
    for i, t in enumerate([
            f"Application = first segment of the Custom name in the '{name}' sheet.",
            f"Point-in-time snapshot of {nrows} line items; not linked to the BoQ sheet.",
            "Confirm: pricing term, Azure Hybrid Benefit, and any inferred storage sizes.",
            "Excludes VAT, bandwidth, backup, ExpressRoute and migration services.",
            "Assumes 730 hours/month of runtime per VM."]):
        put(f"A{n+1+i}", "• " + t, size=9, color="555555", border=False)

    for col, w in zip("ABCDEFG", [34, 10, 10, 26, 17, 17, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{T1+1}"
    ws.sheet_view.showGridLines = False
    wb.active = 0
    wb.save(out)

    print(f"Saved: {out}\n")
    print(f"{'APPLICATION':<26}{'ITEMS':>6}{'vCPU':>7}{'MONTHLY':>14}{'ANNUAL':>15}{'SHARE':>8}")
    print("-" * 76)
    for app, d in order:
        print(f"{app:<26}{d['n']:>6}{d['vcpu']:>7}{d['cost']:>14,.2f}"
              f"{d['cost']*12:>15,.2f}{d['cost']/total:>8.1%}")
    print("-" * 76)
    print(f"{'TOTAL':<26}{nrows:>6}{sum(d['vcpu'] for _, d in order):>7}"
          f"{disp_month:>14,.2f}{disp_year:>15,.2f}{1:>8.1%}")


if __name__ == "__main__":
    main()
