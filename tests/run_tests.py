#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Test suite for the Azure BoQ Creator Skill scripts.

    python tests/run_tests.py

Runs against synthetic fixtures only. No network, no browser, no customer data.

What is covered here:
  - parse_inventory.py   header matching, storage parsing, OS inference,
                         peer-median inference, sheet exclusion
  - size_from_source.py  rightsize vs lift, SKU fitting, disk tiering
  - build_summary.py     Total-row detection, unnamed shared services,
                         category rollup, arithmetic
  - harness.js           syntax only (needs a live page to run properly)

What is NOT covered, and why:
  The browser harness drives a live third-party page whose DOM and prices
  change. Asserting on it would produce a suite that fails for reasons that
  have nothing to do with this code. The harness is instead documented in
  reference/calculator-dom.md, where each failure mode records the observed
  behaviour that justified the workaround.
"""
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPTS = os.path.join(ROOT, "azure-boq-creator-skill", "scripts")
FIXTURES = os.path.join(HERE, "fixtures")

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"  {'PASS' if cond else 'FAIL'}  {name}" + (f"   {detail}" if detail and not cond else ""))


def run(script, *args):
    r = subprocess.run([sys.executable, os.path.join(SCRIPTS, script), *args],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    return r.returncode, (r.stdout or "") + (r.stderr or "")


# ------------------------------------------------------------------ fixtures
def ensure_fixtures():
    needed = ["sample-target-list.xlsx", "sample-source-hardware.xlsx",
              "sample-exported-estimate.xlsx"]
    if all(os.path.exists(os.path.join(FIXTURES, f)) for f in needed):
        return
    print("generating fixtures...")
    subprocess.run([sys.executable, os.path.join(HERE, "make_fixtures.py")], check=True)


# ------------------------------------------------------------ parse_inventory
def test_parse_inventory(tmp):
    print("\nparse_inventory.py")
    out = os.path.join(tmp, "inv.csv")
    code, log = run("parse_inventory.py", os.path.join(FIXTURES, "sample-target-list.xlsx"),
                    "--out", out, "--exclude", "Platform")

    check("exits cleanly", code == 0, log[-400:])
    if code != 0:
        return
    rows = list(csv.DictReader(open(out, encoding="utf-8")))

    check("excluded sheet is skipped", all(r["sheet"] != "Platform" for r in rows))
    check("reads both remaining sheets", {r["sheet"] for r in rows} == {"Orders", "Billing"},
          str({r["sheet"] for r in rows}))
    check("finds every server", len(rows) == 9, f"got {len(rows)}")

    by = {r["vm"]: r for r in rows}

    # storage strings the parser has to survive
    check("parses non-breaking-space '512GB'", by["AZWORDAPP01"]["storage_gb"] == "512")
    check("parses '2 TB - Premium SSD Managed Disks'", by["AZLORDDB01"]["storage_gb"] == "2048")
    check("reads disk type from the text", by["AZLORDDB01"]["disk_type"] == "premiumssd")
    check("reads disk type from a media column", by["AZWBILAPP01"]["disk_type"] == "premiumssd")

    # '?' must inherit from peers in the same application, never a fixed default
    dev = by["AZWORDAPD01"]
    check("'?' storage is inferred, not blank", dev["storage_gb"] not in ("", "0", None))
    check("'?' inference is reported as such", "peer" in dev["storage_src"] or dev["storage_src"] == "default",
          dev["storage_src"])
    peers = sorted(int(r["storage_gb"]) for r in rows if r["sheet"] == "Orders" and r["storage_src"] == "sheet")
    check("'?' inherits the peer median", int(dev["storage_gb"]) == peers[len(peers) // 2],
          f"{dev['storage_gb']} vs median {peers[len(peers)//2]}")

    # OS: explicit column wins, naming convention fills the gap
    check("OS read from column (Linux)", by["AZLORDDB01"]["os_final"] == "linux")
    check("OS read from column (Windows)", by["AZWORDAPP01"]["os_final"] == "windows")
    check("OS inferred when column absent", by["AZWBILAPP01"]["os_final"] == "windows")

    # environment
    check("PROD detected", by["AZWORDAPP01"]["is_prod"] == "True")
    check("QA is not PROD", by["AZWORDAPQ01"]["is_prod"] == "False")
    check("non-PROD defaults to Standard SSD", by["AZWBILAPT01"]["disk_type"] == "standardssd")

    # disk tiering rounds UP, never down - undersizing a disk is a real error
    check("512 GB -> P20 tier", by["AZWORDAPP01"]["disk_code"] == "p20", by["AZWORDAPP01"]["disk_code"])
    check("460 GB rounds up to 512", by["AZWBILAPP01"]["disk_cap_gib"] == "512")
    for r in rows:
        if int(r["disk_cap_gib"]) < int(r["storage_gb"]):
            check("disk tier never undersizes", False, f"{r['vm']}: {r['disk_cap_gib']} < {r['storage_gb']}")
            break
    else:
        check("disk tier never undersizes", True)

    check("label is traceable to the server", all("|" in r["label"] and r["vm"] in r["label"]
                                                  for r in rows if r["vm"]))
    check("inferred values are reported to the user", "CONFIRM" in log or "inferred" in log)


def test_parse_rejects_source_hardware():
    print("\nparse_inventory.py on a source-hardware workbook")
    code, log = run("parse_inventory.py", os.path.join(FIXTURES, "sample-source-hardware.xlsx"),
                    "--out", os.devnull)
    # It must refuse rather than invent SKUs - that refusal is the routing signal.
    check("refuses a workbook with no SKU column", code != 0, f"exit {code}")
    check("explains why", "No server rows found" in log, log[-200:])


# ---------------------------------------------------------- size_from_source
def test_size_from_source(tmp):
    print("\nsize_from_source.py")
    csv_out, plan_out = os.path.join(tmp, "sized.csv"), os.path.join(tmp, "plan.json")
    code, log = run("size_from_source.py", os.path.join(FIXTURES, "sample-source-hardware.xlsx"),
                    "--mode", "rightsize", "--headroom", "0.3",
                    "--region", "westeurope", "--csv", csv_out, "--plan", plan_out)

    check("exits cleanly", code == 0, log[-400:])
    if code != 0:
        return
    rows = list(csv.DictReader(open(csv_out, encoding="utf-8")))

    check("finds every server", len(rows) == 7, f"got {len(rows)}")
    check("ignores the trailing Note row", all("note" not in r["server"].lower() for r in rows))

    by = {r["server"]: r for r in rows}
    check("parses '8.79TB'", abs(float(by["lx1analytics01"]["src_disk"]) - 9001) < 5,
          by["lx1analytics01"]["src_disk"])
    check("parses '755 GB'", float(by["lx1analytics01"]["src_ram"]) == 755)
    check("parses '150G' as GB", float(by["lx1analytics01"]["pk_ram"]) == 150)
    check("parses '64 Lcpu'", float(by["lx1appnode01"]["src_cpu"]) == 64)

    check("Windows detected from OS text", by["wn1web01"]["os"] == "windows")
    check("Oracle Linux detected as linux", by["lx1analytics01"]["os"] == "linux")

    # every recommendation must actually fit the target
    for r in rows:
        if float(r["sku_cpu"]) < float(r["tgt_cpu"]) or float(r["sku_ram"]) < float(r["tgt_ram"]):
            check("every SKU meets its target", False, f"{r['server']} {r['sku']}")
            break
    else:
        check("every SKU meets its target", True)

    plan = json.load(open(plan_out, encoding="utf-8"))
    check("plan has one entry per server", len(plan) == len(rows))
    keys = {"l", "s", "o", "r", "n", "b", "t", "c", "dn"}
    check("plan entries match the harness contract", all(keys <= set(p) for p in plan))
    check("plan carries the requested region", all(p["r"] == "westeurope" for p in plan))
    check("warns that F-series has no RI", "no Reserved Instance" in log or "F-series" in log)


def test_rightsize_vs_lift(tmp):
    print("\nsize_from_source.py  rightsize vs lift")
    res = {}
    for mode in ("rightsize", "lift"):
        c = os.path.join(tmp, f"{mode}.csv")
        code, _ = run("size_from_source.py", os.path.join(FIXTURES, "sample-source-hardware.xlsx"),
                      "--mode", mode, "--headroom", "0.3", "--csv", c,
                      "--plan", os.path.join(tmp, f"{mode}.json"))
        if code != 0:
            check(f"{mode} mode runs", False)
            return
        res[mode] = {r["server"]: r for r in csv.DictReader(open(c, encoding="utf-8"))}

    check("both modes run", len(res) == 2)
    rs, lf = res["rightsize"]["lx1analytics01"], res["lift"]["lx1analytics01"]
    check("rightsize sizes on observed peak", rs["basis"] == "peak")
    check("lift sizes on provisioned spec", lf["basis"] == "provisioned")
    check("over-provisioned server is materially smaller when right-sized",
          float(rs["sku_ram"]) < float(lf["sku_ram"]),
          f"{rs['sku']} ({rs['sku_ram']}GB) vs {lf['sku']} ({lf['sku_ram']}GB)")

    tot_rs = sum(float(r["sku_cpu"]) for r in res["rightsize"].values())
    tot_lf = sum(float(r["sku_cpu"]) for r in res["lift"].values())
    check("right-sizing reduces total vCPU", tot_rs < tot_lf, f"{tot_rs} vs {tot_lf}")


# ------------------------------------------------------------- build_summary
def test_build_summary(tmp):
    print("\nbuild_summary.py")
    try:
        import openpyxl
    except ImportError:
        check("openpyxl available", False, "pip install openpyxl")
        return

    out = os.path.join(tmp, "boq.xlsx")
    code, log = run("build_summary.py", os.path.join(FIXTURES, "sample-exported-estimate.xlsx"),
                    "--out", out, "--title", "Test BoQ")
    check("exits cleanly", code == 0, log[-400:])
    if code != 0:
        return

    wb = openpyxl.load_workbook(out)
    check("adds a Summary sheet", "Summary" in wb.sheetnames)
    check("keeps the original sheet", len(wb.sheetnames) >= 2)
    check("Summary is the active sheet", wb.active.title == "Summary")

    ws = wb["Summary"]
    check("counts every priced line item", ws["B4"].value == 7, f"got {ws['B4'].value}")

    EXPECTED = 444.33 + 444.33 + 972.09 + 502.45 + 21.50 + 87.04 + 10327.00
    check("monthly total matches the export", abs(ws["B5"].value - EXPECTED) < 0.05,
          f"{ws['B5'].value} vs {EXPECTED}")
    check("annual is 12x monthly", abs(ws["B6"].value - ws["B5"].value * 12) < 0.5)
    check("3-year is 36x monthly", abs(ws["B7"].value - ws["B5"].value * 36) < 1.0)

    # walk the application table
    apps, r = {}, 11
    while ws[f"A{r}"].value and str(ws[f"A{r}"].value) != "TOTAL":
        apps[str(ws[f"A{r}"].value)] = (ws[f"B{r}"].value, ws[f"E{r}"].value)
        r += 1

    check("groups by application", "Orders" in apps and "Billing" in apps, str(list(apps)))
    check("Orders has 3 items", apps.get("Orders", (0,))[0] == 3)
    # a shared service with no Custom name must fall back to its service type,
    # not be dropped - this was a real bug
    check("unnamed shared service is not dropped", "Storage Accounts" in apps, str(list(apps)))
    check("application shares sum to the total",
          abs(sum(v[1] for v in apps.values()) - ws["B5"].value) < 0.05)
    check("largest application is listed first",
          max(apps.values(), key=lambda v: v[1])[1] == list(apps.values())[0][1])

    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    for table in ("COST BY APPLICATION", "COST BY SERVICE CATEGORY",
                  "COST BY REGION", "COST BY PRICING MODEL"):
        check(f"has a {table.lower()} table", table in text)

    check("distinguishes reserved from savings plan",
          "3-Year Reserved Instance" in text and "3-Year Savings Plan" in text)
    check("states its assumptions", "assumption" in text.lower() or "Confirm" in text)

    # cross-sheet formulas render as #ERROR! in the app's preview, so the
    # figures a reader sees must be static values
    check("headline figures are static, not cross-sheet formulas",
          not (isinstance(ws["B5"].value, str) and str(ws["B5"].value).startswith("=")))


def test_summary_is_idempotent(tmp):
    print("\nbuild_summary.py  re-running on its own output")
    try:
        import openpyxl
    except ImportError:
        return
    once = os.path.join(tmp, "once.xlsx")
    twice = os.path.join(tmp, "twice.xlsx")
    run("build_summary.py", os.path.join(FIXTURES, "sample-exported-estimate.xlsx"), "--out", once)
    code, log = run("build_summary.py", once, "--out", twice)
    check("re-runs without error", code == 0, log[-300:])
    if code != 0:
        return
    a, b = openpyxl.load_workbook(once)["Summary"], openpyxl.load_workbook(twice)["Summary"]
    check("produces the same total", abs((a["B5"].value or 0) - (b["B5"].value or 0)) < 0.05)
    check("does not duplicate the Summary sheet",
          openpyxl.load_workbook(twice).sheetnames.count("Summary") == 1)


# ------------------------------------------------------------------ harness
def test_harness_syntax():
    print("\nharness.js")
    path = os.path.join(ROOT, "azure-boq-creator-skill", "scripts", "harness.js")
    src = open(path, encoding="utf-8").read()

    node = shutil.which("node")
    if node:
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
            f.write("(async function(){\n" + src + "\n})")
            tmpf = f.name
        r = subprocess.run([node, "--check", tmpf], capture_output=True, text=True)
        os.unlink(tmpf)
        check("parses as valid JavaScript", r.returncode == 0, (r.stderr or "")[:300])
    else:
        print("  SKIP  node not on PATH")

    for fn in ["__setNative", "__rows", "__expand", "__collapse", "__addVM", "__addProduct",
               "__findProducts", "__describeRow", "__cfgGeneric", "__cfgRow", "__pickInstance",
               "__setBilling", "__setDisk", "runPlan", "runMixedPlan", "__auditRows",
               "__auditCosts", "__quickAudit", "__totals", "__export"]:
        check(f"exports {fn}", f"window.{fn}" in src)

    # guards against regressions of specific, verified bugs
    check("filters the phantom empty row",
          "button.module-name" in src and "wa-calcService" in src)
    check("audits for zero-cost line items", "ZERO_COST" in src)
    check("audits label quantity against configuration", "LABEL_QTY_NOT_IN_ROW" in src)
    check("audits for default regions", "REGION_MISMATCH" in src)
    check("falls back to a savings plan when no RI exists", "sv-three-year" in src)
    check("reads product titles from the SVG title, not innerText",
          "data-slug-id" in src and "title" in src)


# --------------------------------------------------------------------- docs
def test_docs():
    print("\ndocumentation")
    base = os.path.join(ROOT, "azure-boq-creator-skill")
    skill = open(os.path.join(base, "SKILL.md"), encoding="utf-8").read()

    check("SKILL.md has frontmatter", skill.startswith("---"))
    m = re.search(r"^name:\s*(\S+)", skill, re.M)
    check("frontmatter declares the skill name", bool(m) and m.group(1) == "azure-boq-creator-skill")
    d = re.search(r"^description:\s*(.+)$", skill, re.M)
    check("frontmatter has a description", bool(d) and len(d.group(1)) > 80)

    for ref in ["calculator-dom.md", "products.md", "service-catalogue.md", "fabric-sizing.md"]:
        check(f"reference/{ref} exists", os.path.exists(os.path.join(base, "reference", ref)))
        check(f"SKILL.md points at {ref}", ref in skill)

    for s in ["harness.js", "parse_inventory.py", "size_from_source.py", "build_summary.py"]:
        check(f"SKILL.md documents {s}", s in skill)

    dom = open(os.path.join(base, "reference", "calculator-dom.md"), encoding="utf-8").read()
    modes = re.findall(r"^### (\d+)\.", dom, re.M)
    check("failure modes are numbered contiguously",
          [int(x) for x in modes] == list(range(1, len(modes) + 1)), str(modes))
    check("SKILL.md's failure-mode count is accurate", f"{len(modes)} failure modes" in skill,
          f"doc has {len(modes)}")

    cat = open(os.path.join(base, "reference", "service-catalogue.md"), encoding="utf-8").read()
    for name in ["Foundry IQ", "Foundry Tools", "Virtual Machine Scale Sets",
                 "Azure Elastic SAN", "Microsoft Sentinel", "Azure OpenAI", "Microsoft Fabric"]:
        check(f"catalogue covers {name}", name in cat)
    check("catalogue lists what is NOT priceable", "not on the calculator" in cat.lower())

    fab = open(os.path.join(base, "reference", "fabric-sizing.md"), encoding="utf-8").read()
    check("fabric doc points at the estimator", "estimator.fabric.microsoft.com" in fab)
    check("fabric doc maps F256 to its calculator value", "twofiftysix" in fab)
    check("fabric doc warns the hours factor multiplies", "MULTIPLIES" in fab or "multiplies" in fab)


# ---------------------------------------------------------------- no secrets
def test_no_customer_data():
    print("\nrepository hygiene")
    # Patterns from the real engagements this was built against. None of them
    # may ever appear in this repository. Assembled from fragments so this
    # scanner does not match its own pattern list.
    banned = [
        ("QGRL" + r"[A-Z0-9]+", "customer hostname"),
        ("AZ" + "QW" + r"[A-Z0-9]{4,}", "customer hostname"),
        ("D1" + "W" + r"[A-Z0-9]{4,}", "customer hostname"),
        ("li1" + "qnbfs" + "|" + "sas" + "qnbfs" + "|" + "WF1" + "BR", "customer hostname"),
        (r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", "GUID"),
        ("gh" + r"[pousr]_[A-Za-z0-9]{20,}", "credential"),
    ]
    ALLOWED_GUID = "00000000-0000-0000-0000-0000000000"  # obvious placeholders
    SELF = os.path.basename(__file__)

    offenders = []
    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = [d for d in dirnames if d not in (".git", "__pycache__")]
        for fn in filenames:
            if fn.endswith((".xlsx", ".png", ".jpg", ".pyc")) or fn == SELF:
                continue
            p = os.path.join(dirpath, fn)
            try:
                txt = open(p, encoding="utf-8", errors="ignore").read()
            except OSError:
                continue
            for pat, what in banned:
                for hit in re.findall(pat, txt):
                    if what == "GUID" and hit.startswith(ALLOWED_GUID):
                        continue
                    offenders.append(f"{os.path.relpath(p, ROOT)}: {what} '{hit[:24]}'")
    check("no customer hostnames, GUIDs or credentials in the repo",
          not offenders, "; ".join(offenders[:4]))

    # the fixtures must also be free of real data
    fx = os.path.join(FIXTURES, "sample-target-list.xlsx")
    if os.path.exists(fx):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(fx)
            cells = [str(c.value) for s in wb.worksheets for row in s.iter_rows()
                     for c in row if c.value is not None]
            bad = [c for c in cells
                   if re.search("QGRL" + r"[A-Z0-9]+", c) or re.search("AZ" + "QW" + r"[A-Z0-9]{4,}", c)]
            check("fixture workbooks contain no real hostnames", not bad, str(bad[:3]))
        except ImportError:
            pass


# --------------------------------------------------------------------- main
def main():
    ensure_fixtures()
    tmp = tempfile.mkdtemp(prefix="boq-tests-")
    try:
        test_parse_inventory(tmp)
        test_parse_rejects_source_hardware()
        test_size_from_source(tmp)
        test_rightsize_vs_lift(tmp)
        test_build_summary(tmp)
        test_summary_is_idempotent(tmp)
        test_harness_syntax()
        test_docs()
        test_no_customer_data()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    total = len(PASS) + len(FAIL)
    print(f"\n{'-' * 62}\n{len(PASS)}/{total} passed")
    if FAIL:
        print("\nfailed:")
        for f in FAIL:
            print(f"  - {f}")
        return 1
    print("all green")
    return 0


if __name__ == "__main__":
    sys.exit(main())
