# -*- coding: utf-8 -*-
"""
Build the example inventory used in TUTORIAL.md.

Fictional company "Northwind Retail". Every hostname, application and figure is
invented - safe to show on camera and to share publicly.

Designed so the skill's reasoning is *visible* in a short demo:
  - 10 servers in scope, so a full build takes ~90s
  - two sheets with completely different column names  -> fuzzy header matching
  - storage written as 1.5TB / 600GB / "2 TB - Premium SSD Managed Disks" / ?
  - one sheet with no OS column at all                 -> naming-convention inference
  - three '?' storage cells                            -> peer-median inference
  - a container platform sheet                          -> the "what's in scope" question
  - PROD and non-PROD mixed                            -> the RI vs PAYG question
  - an F-series PROD server                            -> the "no Reserved Instance" fallback
"""
import openpyxl

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Northwind-Retail-Server-Inventory.xlsx")
wb = openpyxl.Workbook()

# ---- Sheet 1: Storefront -------------------------------------------------
ws = wb.active
ws.title = "Storefront"
ws.append(["Application name", "Role", "Environment", "Server Name at Source",
           "CPU at Source", "RAM at Source", "Server Name at Target", "OS Version",
           "VM Type at Target", "Storage Disks at Target"])
for r in [
    ("Storefront", "Web", "PROD", "NWSRCWEB01", 8, 32, "NWAZWWEB01", "Windows Server 2022",
     "Standard_D8s_v5", "\xa0512GB"),
    ("Storefront", "Web", "PROD", "NWSRCWEB02", 8, 32, "NWAZWWEB02", "Windows Server 2022",
     "Standard_D8s_v5", " 512GB"),
    ("Storefront", "DB", "PROD", "NWSRCDB01", 16, 128, "NWAZLDB01", "RHEL 9",
     "Standard_E16s_v5", "2 TB - Premium SSD Managed Disks"),
    ("Storefront", "Search", "PROD", "NWSRCSCH01", 8, 16, "NWAZWSCH01", "Windows Server 2022",
     "Standard_F8s_v2", "1.5TB"),
    ("Storefront", "Web", "QA", "NWSRCWEBQ01", 8, 32, "NWAZWWEBQ01", "Windows Server 2022",
     "Standard_D8s_v5", "600GB"),
    ("Storefront", "DB", "QA", "NWSRCDBQ01", 8, 64, "NWAZLDBQ01", "RHEL 9",
     "Standard_E8s_v5", "?"),
]:
    ws.append(list(r))

# ---- Sheet 2: Warehouse (different headers, NO OS column) ----------------
ws2 = wb.create_sheet("Warehouse")
ws2.append(["Server Role", "Server Name", "Server Name (Actual)",
            "Environment(DEV/QA/PROD)", "CPU", "RAM [GB]", "Storage Size [GB]",
            "Azure VM SKU", "Storage Media"])
for r in [
    ("App", "Warehouse Mgmt", "NWAZWWMS01", "PROD", 8, 64, 800, "Standard_E8s_v5", "PremiumSSD"),
    ("App", "Warehouse Mgmt", "NWAZWWMS02", "PROD", 8, 64, 800, "Standard_E8s_v5", "PremiumSSD"),
    ("App", "Warehouse Mgmt", "NWAZWWMSD01", "DEV", 4, 16, "?", "Standard_D4s_v5", "StandardSSD"),
    ("Reporting", "Warehouse Mgmt", "NWAZWRPT01", "QA", 4, 16, "?", "Standard_D4s_v5", "StandardSSD"),
]:
    ws2.append(list(r))

# ---- Sheet 3: container platform - prompts the scope question ------------
ws3 = wb.create_sheet("Container Platform")
ws3.append(["Cluster #", "VM", "CPU", "Mem", "Workload Env", "Zone1", "SKU Recommendation"])
for r in [
    ("Cluster1", "control-plane-1", "8vCPU", "32GB", "MGMT", "AZ1", "Standard_D8s_v5"),
    ("Cluster1", "control-plane-2", "8vCPU", "32GB", "MGMT", "AZ2", "Standard_D8s_v5"),
    ("Cluster1", "control-plane-3", "8vCPU", "32GB", "MGMT", "AZ3", "Standard_D8s_v5"),
    ("Cluster1", "worker-1", "16vCPU", "128GB", "Internal STG/PRD", "AZ1", "Standard_E16s_v5"),
    ("Cluster1", "worker-2", "16vCPU", "128GB", "Internal STG/PRD", "AZ2", "Standard_E16s_v5"),
]:
    ws3.append(list(r))

for sheet in wb.worksheets:
    for col, w in zip("ABCDEFGHIJ", [20, 14, 14, 22, 10, 12, 22, 22, 22, 30]):
        sheet.column_dimensions[col].width = w
    for c in sheet[1]:
        c.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        c.fill = openpyxl.styles.PatternFill("solid", fgColor="2E5B9A")
    sheet.freeze_panes = "A2"

wb.save(OUT)
print("wrote", OUT)

wb2 = openpyxl.load_workbook(OUT)
total = 0
for s in wb2.worksheets:
    n = sum(1 for r in s.iter_rows(min_row=2, values_only=True) if r[0])
    total += n
    print(f"  {s.title:<20} {n} servers")
print(f"  {'TOTAL':<20} {total}  ({total - 5} in scope if Container Platform is excluded)")
